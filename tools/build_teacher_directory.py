"""Build the supplied, masked teacher introductions and exact branch links.

The public release contains rendered introductions and representative photos,
never the supplied workbook or private inspection records.
"""
from pathlib import Path
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import quote, unquote, urlsplit
from html import escape, unescape
import argparse, hashlib, json, re, shutil, struct, zipfile
import openpyxl
from lxml import html

ROOT=Path(__file__).resolve().parents[1]
DOMAIN='https://xn--3e0bl59bm0ad17a.com'
DAY='2026-10-02'
E=lambda value:escape(str(value),quote=True)
J=lambda obj:json.dumps(obj,ensure_ascii=False,separators=(',',':')).replace('<','\\u003c')
H=lambda path:quote(path,safe='/#-._~')
digest=lambda data:hashlib.sha256(data).hexdigest()
MENU='<a href="/선생님찾기/" data-teacher-menu>선생님찾기</a>'
# Source names are preserved in each introduction. These aliases only join the
# existing directory: 모두/석사 are confirmed by the center workbook; 주엽2호
# is explicitly confirmed as the same location by the user on 2026-10-02.
ALIASES={'목감점(모두)':'목감점','별내중앙점(모두)':'별내중앙점','용인백현점(모두)':'용인백현점',
 '진천점(모두)':'진천점','탕정점(모두)':'탕정점','석사2호점':'석사점','주엽2호점':'주엽점'}
REGIONS=['서울','경기','인천','강원','대전','세종','충북','충남','대구','부산','울산','경북','경남','광주','전북','전남','제주']
REGION_NAMES={'서울특별시':'서울','경기도':'경기','인천광역시':'인천','강원특별자치도':'강원','강원도':'강원',
 '대전광역시':'대전','세종특별자치시':'세종','충청북도':'충북','충청남도':'충남','대구광역시':'대구',
 '부산광역시':'부산','울산광역시':'울산','경상북도':'경북','경상남도':'경남','광주광역시':'광주','전북특별자치도':'전북','전라북도':'전북','전라남도':'전남','제주특별자치도':'제주'}

def write(path,text):
    path.parent.mkdir(parents=True,exist_ok=True)
    data=(text.rstrip()+'\n').encode('utf-8')
    if not path.exists() or path.read_bytes()!=data:path.write_bytes(data)

def load(path):return json.loads(path.read_text('utf-8'))
def route(branch=None):return '/선생님찾기/'+(branch['slug']+'/' if branch else '')
def norm(value):return re.sub(r'\s+','',str(value or ''))

def read_model(teacher_dir,center_file):
    source=teacher_dir/'교사 프로필.xlsx'
    wb=openpyxl.load_workbook(source,read_only=True,data_only=True)
    rows=[]
    for sheet in wb:
        for row,values in enumerate(sheet.iter_rows(values_only=True),1):
            if not any(value is not None for value in values):continue
            assert len(values)==4 and all(isinstance(v,str) and v.strip() for v in values),(sheet.title,row,'Expected four text cells')
            branch,name,focus,bio=[v.strip() for v in values]
            assert '*' in name,(row,'Preserve source-masked names')
            assert bio.startswith(branch+' '+name+' 선생님입니다.'),(row,'Introduction header differs')
            rows.append({'sourceSheet':sheet.title,'sourceRow':row,'sourceBranch':branch,'name':name,'focus':focus,'bio':bio,'id':f'teacher-{row:04d}'})
    assert len(rows)==1002
    centers={}
    center_wb=openpyxl.load_workbook(center_file,read_only=True,data_only=True)
    for sheet in center_wb:
        for row,values in enumerate(sheet.iter_rows(values_only=True),1):
            if row==1 or not values[0]:continue
            centers[str(values[0]).strip()]={'address':' '.join(str(values[11] or '').split()),'officialName':' '.join(str(values[6] or '').split()),'sourceRow':row}
    site={}
    for path in sorted((ROOT/'지점안내').glob('*/*/index.html')):
        doc=html.fromstring(path.read_bytes())
        base='/'+path.relative_to(ROOT).as_posix().removesuffix('index.html')
        neighborhood_links=[]
        for a in doc.xpath('//*[@id="related-pages"]//a[@href]'):
            u=unquote(urlsplit(a.get('href')).path)
            if u.startswith('/전국학원/') and len(u.strip('/').split('/'))==4:
                neighborhood_links.append({'path':u,'name':u.strip('/').split('/')[-1]})
        address=doc.xpath('string(//*[@id="center-facts"]//dt[normalize-space()="주소"]/following-sibling::dd[1])').strip()
        assert address,(path,'Existing branch address missing')
        site[path.parent.name]={'path':base,'region':path.parent.parent.name,'address':address,'neighborhoods':neighborhood_links}
    grouped=defaultdict(list)
    for row in rows:grouped[ALIASES.get(row['sourceBranch'],row['sourceBranch'])].append(row)
    photos=[]
    for p in sorted(teacher_dir.glob('collage-profile-*.png')):
        # Exclude two very similar compositions from the supplied contact sheet.
        if p.name in ('collage-profile-06.png','collage-profile-09.png'):continue
        data=p.read_bytes(); width,height=struct.unpack('>II',data[16:24])
        photos.append({'source':p.name,'path':'assets/teacher-profiles/intro-'+p.stem.rsplit('-',1)[-1]+'.png','width':width,'height':height,'sha256':digest(data)})
    assert len(photos)==10 and len({p['sha256'] for p in photos})==10
    branches=[]
    for slug,teachers in grouped.items():
        raw_names=list(dict.fromkeys(t['sourceBranch'] for t in teachers))
        source_name=next((n for n in raw_names if n in centers),slug)
        center=centers.get(source_name) or centers.get(slug)
        if slug in site:
            center={'address':site[slug]['address'],'sourceRow':center['sourceRow'] if center else None}
        assert center,(slug,'Unresolved center source')
        address_region=center['address'].split()[0]
        region=site.get(slug,{}).get('region') or REGION_NAMES.get(address_region,address_region)
        assert region in REGIONS,(slug,region)
        assert len(teachers)<=len(photos),(slug,'Not enough distinct images')
        shift=int(digest(slug.encode())[:8],16)%len(photos)
        for index,t in enumerate(teachers):t['photo']=photos[(shift+index)%len(photos)]
        branches.append({'slug':slug,'name':slug,'sourceNames':raw_names,'region':region,'address':center['address'],
          'centerSourceName':source_name,'centerSourceRow':center['sourceRow'],'sitePath':site.get(slug,{}).get('path'),
          'neighborhoods':site.get(slug,{}).get('neighborhoods',[]),'teachers':teachers})
    branches.sort(key=lambda b:(REGIONS.index(b['region']),b['name']))
    assert len(branches)==204 and len([b for b in branches if b['sitePath']])==193
    return {'sourceHash':digest(source.read_bytes()),'centerSourceHash':digest(center_file.read_bytes()),
      'photoUse':'User confirmed representative shared images, 2026-10-02','sourceRows':len(rows),'sourceBranchLabels':len(grouped)+1,
      'aliases':ALIASES,'branches':branches,'photos':photos}

def add_menu(source):
    for pattern in [r'(<div\b[^>]*data-unified-nav[^>]*>)([\s\S]*?)(</div>)',r'(<nav\b[^>]*class="footer-links"[^>]*>)([\s\S]*?)(</nav>)']:
        def replace(match):
            if 'data-teacher-menu' in match[2]:return match[0]
            links=re.sub(r'(<a\b[^>]*>지점안내</a>)',r'\1'+MENU,match[2],count=1)
            assert links!=match[2],'Expected existing branch navigation'
            return match[1]+links+match[3]
        source=re.sub(pattern,replace,source,count=1)
    source=re.sub(r'(/assets/unified-ui\.css)(?:\?[^"\s>]*)?',r'\1?v=20261002-teachers',source)
    return source

def shell(path,title,desc,body,header,footer,items):
    assert 0<len(desc)<=80 and desc.endswith('.')
    current=header.replace(' class="active" aria-current="page"','')
    current=current.replace('data-teacher-menu>','data-teacher-menu class="active" aria-current="page">')
    url=DOMAIN+H(path)
    crumbs=[('홈','/'),('선생님찾기','/선생님찾기/')]+([(title.removesuffix(' 선생님 소개'),path)] if path!='/선생님찾기/' else [])
    crumb_html='<nav class="td-breadcrumb" aria-label="현재 위치">'+'<span aria-hidden="true"> / </span>'.join(f'<a href="{H(p)}">{E(n)}</a>' if i<len(crumbs)-1 else f'<span aria-current="page">{E(n)}</span>' for i,(n,p) in enumerate(crumbs))+'</nav>'
    graph=[{'@type':'Organization','@id':DOMAIN+'/#organization','name':'전국학원','url':DOMAIN+'/'},
      {'@type':'CollectionPage','@id':url+'#webpage','url':url,'name':title,'description':desc,'inLanguage':'ko-KR','datePublished':DAY,'dateModified':DAY,'publisher':{'@id':DOMAIN+'/#organization'},'breadcrumb':{'@id':url+'#breadcrumb'},'mainEntity':{'@id':url+'#list'}},
      {'@type':'BreadcrumbList','@id':url+'#breadcrumb','itemListElement':[{'@type':'ListItem','position':i,'name':n,'item':DOMAIN+H(p)} for i,(n,p) in enumerate(crumbs,1)]},
      {'@type':'ItemList','@id':url+'#list','numberOfItems':len(items),'itemListElement':[{'@type':'ListItem','position':i,'name':n,'url':DOMAIN+H(p)} for i,(n,p) in enumerate(items,1)]}]
    return f'''<!doctype html>
<html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{E(title)} | 전국학원</title><meta name="description" content="{E(desc)}"><meta name="robots" content="index,follow"><link rel="canonical" href="{url}"><link rel="icon" href="/assets/favicon.png"><meta name="theme-color" content="#173b32">
<meta property="og:type" content="website"><meta property="og:locale" content="ko_KR"><meta property="og:site_name" content="전국학원"><meta property="og:title" content="{E(title)} | 전국학원"><meta property="og:description" content="{E(desc)}"><meta property="og:url" content="{url}"><meta property="og:image" content="{DOMAIN}/assets/title.png"><meta name="twitter:card" content="summary_large_image"><meta name="twitter:title" content="{E(title)} | 전국학원"><meta name="twitter:description" content="{E(desc)}"><meta name="twitter:image" content="{DOMAIN}/assets/title.png">
<link rel="stylesheet" href="/assets/site.css"><link rel="stylesheet" href="/assets/unified-ui.css?v=20261002-teachers"><link rel="stylesheet" href="/assets/teacher-directory.css?v=20261002"><script defer src="/assets/unified-ui.js?v=20260914-2"></script><script defer src="/assets/teacher-directory.js?v=20261002"></script><script type="application/ld+json">{J({'@context':'https://schema.org','@graph':graph})}</script></head>
<body class="unified-ui td-page" data-teacher-directory="2026-10-02"><a class="skip-link" href="#main">본문 바로가기</a>{current}<main id="main"><div class="wrap">{crumb_html}{body}</div></main>{footer}</body></html>'''

def branch_card(b):
    names=' · '.join(t['name'] for t in b['teachers'])
    focuses=list(dict.fromkeys(t['focus'].split(' / ')[0] for t in b['teachers']))
    search=' '.join([b['name'],*b['sourceNames'],b['region'],b['address'],names,*focuses,*[x['name'] for x in b['neighborhoods']]])
    return f'''<article class="td-branch-card" data-branch-card data-region="{E(b['region'])}" data-search="{E(search)}" data-focus="{E(' | '.join(focuses))}" data-count="{len(b['teachers'])}"><div class="td-card-top"><span>{E(b['region'])}</span><span>소개 {len(b['teachers'])}건</span></div><h3><a href="{H(route(b))}">{E(b['name'])}<span aria-hidden="true"> ↗</span></a></h3><p class="td-place">{E(b['address'])}</p><p class="td-names">{E(names)}</p><p class="td-focus-preview">{' · '.join(E(f) for f in focuses[:2])}</p></article>'''

def hub(model,header,footer):
    branches=model['branches']; focuses=sorted({t['focus'].split(' / ')[0] for b in branches for t in b['teachers']})
    region_options=''.join(f'<option value="{E(r)}">{E(r)}</option>' for r in REGIONS if any(b['region']==r for b in branches))
    focus_options=''.join(f'<option value="{E(f)}">{E(f)}</option>' for f in focuses)
    body=f'''<section class="td-hero"><div><p class="td-eyebrow">TEACHER DIRECTORY</p><h1>선생님찾기</h1><p class="td-lead">학생을 어떻게 이해하고, 학습을 어떻게 도울까요?<br>가까운 지점 선생님의 소개에서 살펴보세요.</p><div class="td-stats"><span><b>{len(branches)}</b>개 지점</span><span><b>{model['sourceRows']:,}</b>건의 교사 소개</span></div><a class="btn btn-primary" href="#find-teachers">우리 지점 선생님 찾아보기 ↓</a></div><aside class="td-intro"><strong>소개에서 이런 내용을 살펴보세요</strong><ol><li><b>학습 지도 방향</b><span>질문·오답·복습 등 선생님이 중요하게 여기는 과정</span></li><li><b>선생님의 인사</b><span>학생의 학습을 함께 살펴보는 방법</span></li><li><b>상담으로 이어가기</b><span>학생의 현재 과제와 궁금한 점을 준비하세요.</span></li></ol></aside></section>
<section id="find-teachers" class="td-directory"><p class="td-eyebrow">FIND YOUR CENTER</p><h2>지점별 선생님 소개</h2><p>지점을 선택하면 선생님별 소개와 학습 지도 방향을 함께 볼 수 있습니다.</p>
<div class="td-filter" data-teacher-filter hidden><div class="td-filter-row"><div><label for="teacher-region">지역</label><select id="teacher-region"><option value="all">모든 지역</option>{region_options}</select></div><div><label for="teacher-focus">학습 지도 방향</label><select id="teacher-focus"><option value="all">모든 지도 방향</option>{focus_options}</select></div><div class="td-search"><label for="teacher-search">지점·동네·선생님 검색</label><input id="teacher-search" type="search" placeholder="예: 명일점, 주엽동, 오답" autocomplete="off"></div><button class="btn" type="button" id="teacher-reset">초기화</button></div><p id="teacher-count" role="status" aria-live="polite"></p></div>
<div class="td-branch-grid" id="teacher-results">{''.join(branch_card(b) for b in branches)}</div><div class="td-empty" id="teacher-empty" hidden><h3>검색 결과가 없습니다</h3><p>지점명이나 검색어를 짧게 입력해 보세요.</p><button class="btn" type="button" data-teacher-reset>전체 지점 보기</button></div><div class="td-more" hidden><button class="btn" type="button" id="teacher-more">지점 더 보기</button></div></section>
<section class="td-bottom"><div><h2>상담 전에 궁금한 점을 정리해 보세요</h2><p>학생의 현재 과제와 어려웠던 부분을 함께 준비하면 필요한 도움을 이야기하기 좋습니다. 실제 담당 과목·학년과 수업 배정은 지점에 문의해 주세요.</p></div><a class="btn" href="{H('/학습가이드/학습진단상담준비/')}">상담 준비 가이드</a></section>'''
    desc='전국 지점별 선생님의 소개와 학습 지도 방향을 확인하고, 가까운 지점의 교사 정보를 찾아보세요.'
    return shell(route(),'선생님찾기',desc,body,header,footer,[(b['name']+' 선생님 소개',route(b)) for b in branches]),desc

def teacher_card(t):
    photo=t['photo']; sentences=re.split(r'(?<=[.!?])\s+',t['bio']); paragraphs=[' '.join(sentences[i:i+2]) for i in range(0,len(sentences),2)]
    tags=''.join(f'<span>{E(f)}</span>' for f in t['focus'].split(' / '))
    return f'''<article class="td-teacher-card" id="{t['id']}" data-teacher-card data-source-row="{t['sourceRow']}"><div class="td-profile-head"><figure><img src="/{photo['path']}" width="{photo['width']}" height="{photo['height']}" alt="교사 소개용 공용 이미지" loading="lazy" decoding="async"><figcaption>소개용 이미지</figcaption></figure><div><p class="td-teacher-branch">{E(t['sourceBranch'])}</p><h2>{E(t['name'])} <span>선생님</span></h2><p class="td-focus-label">학습 지도 방향</p><div class="td-tags">{tags}</div></div></div><div class="td-bio">{''.join('<p>'+E(p)+'</p>' for p in paragraphs)}</div></article>'''

def branch_page(b,model,header,footer):
    teachers=b['teachers']; quick=''.join(f'<a href="#{t["id"]}">{E(t["name"])} 선생님</a>' for t in teachers)
    branch_link=(f'<a class="btn" href="{H(b["sitePath"])}">{E(b["name"])} 위치·수강 정보</a>' if b['sitePath'] else '')
    merged='<p class="td-merge-note">주엽점과 주엽2호점의 선생님 소개를 함께 안내합니다.</p>' if b['slug']=='주엽점' else ''
    neighborhoods=''.join(f'<a href="{H(n["path"])}">{E(n["name"])} 학원 안내</a>' for n in b['neighborhoods'])
    near=[other for other in model['branches'] if other['region']==b['region'] and other['slug']!=b['slug']][:3]
    body=f'''<section class="td-branch-hero"><p class="td-eyebrow">{E(b['region'])} · 선생님찾기</p><h1>{E(b['name'])} 선생님 소개</h1><p class="td-lead">선생님의 인사와 학습 지도 방향을 살펴보세요.</p><p class="td-location">{E(b['address'])}</p>{merged}<div class="td-branch-actions"><a class="btn btn-primary" href="#teacher-profiles">교사 소개 {len(teachers)}건 보기 ↓</a>{branch_link}</div></section>
<section id="teacher-profiles"><div class="td-section-heading"><h2>함께 학습을 살펴보는 선생님</h2><span>소개 {len(teachers)}건</span></div><p class="td-photo-note">사진은 선생님 소개를 위한 공용 이미지이며 실제 교사 사진이 아닙니다.</p><nav class="td-teacher-jump" aria-label="선생님 소개 바로가기">{quick}</nav><div class="td-profile-grid">{''.join(teacher_card(t) for t in teachers)}</div></section>
<section class="td-consult"><p class="td-eyebrow">상담 전에</p><h2>학생에게 필요한 도움을 함께 이야기해 주세요</h2><p>현재 학년과 과목, 최근에 어려웠던 과제를 준비해 주세요. 궁금한 지도 방식은 선생님 소개를 보며 메모해 두면 좋습니다. 실제 담당 과목·학년과 수업 배정, 상담 일정은 지점에 확인해 주세요.</p><div class="td-branch-actions"><a class="btn btn-primary" href="/상담문의/">상담 문의</a><a class="btn" href="{H('/학습가이드/학습진단상담준비/')}">상담 준비 가이드</a></div></section>
{('<section class="td-neighborhoods"><h2>우리 동네 학원 안내</h2><div>'+neighborhoods+'</div></section>') if neighborhoods else ''}
<section class="td-related"><div class="td-section-heading"><h2>{E(b['region'])}의 다른 지점도 살펴보세요</h2><a href="{H(route())}">전체 지점 보기 →</a></div><div class="td-branch-grid">{''.join(branch_card(other).replace('data-branch-card','data-related-branch') for other in near)}</div></section><p class="td-updated">소개 내용 반영 <time datetime="{DAY}">2026.10.02</time></p>'''
    desc=f'{b["name"]} 선생님들의 학습 지도 방향과 소개를 살펴보고, 상담 전에 궁금한 내용을 준비하세요.'
    return shell(route(b),b['name']+' 선생님 소개',desc,body,header,footer,[(t['name']+' 선생님',route(b)+'#'+t['id']) for t in teachers]),desc

def bridge(b):
    return f'''<section class="teacher-bridge" id="teacher-introductions" data-teacher-bridge="{E(b['slug'])}"><div><p class="teacher-bridge-kicker">선생님 소개</p><h2>{E(b['name'])} 선생님을 만나보세요</h2><p>선생님의 인사와 학습 지도 방향을 살펴보고, 상담에서 궁금한 내용을 준비해 보세요.</p></div><a href="{H(route(b))}">교사 소개 {len(b['teachers'])}건 보기 <span aria-hidden="true">↗</span></a></section>'''

def main():
    parser=argparse.ArgumentParser(); parser.add_argument('--teacher-dir',required=True);parser.add_argument('--center-file',required=True);parser.add_argument('--report-dir',required=True);a=parser.parse_args()
    teacher_dir=Path(a.teacher_dir); out=Path(a.report_dir);out.mkdir(parents=True,exist_ok=True)
    model=read_model(teacher_dir,Path(a.center_file));manifest=load(ROOT/'release-public-manifest.json')
    baseline_path=out/'before-release-public-manifest.json'
    if not baseline_path.exists():shutil.copy2(ROOT/'release-public-manifest.json',baseline_path)
    baseline=load(baseline_path)
    for filename in ['seo-descriptions.json','sitemap.xml']:
        if not (out/('before-'+filename)).exists():shutil.copy2(ROOT/filename,out/('before-'+filename))
    backup=out/'before-pages.zip'
    if not backup.exists():
        with zipfile.ZipFile(backup,'w',zipfile.ZIP_DEFLATED,compresslevel=1) as z:
            for name in baseline['files']:
                if name.endswith('.html'):z.write(ROOT/name,name)
    template=add_menu((ROOT/'학습가이드/index.html').read_text('utf-8'))
    header=re.search(r'<header class="site-header">[\s\S]*?</header>',template)[0]
    footer=re.search(r'<footer class="site-footer">[\s\S]*?</footer>',template)[0]
    changed=set();new_pages=[];context_links={};nav_count=0
    by_slug={b['slug']:b for b in model['branches']}; by_site={b['sitePath']:b for b in model['branches'] if b['sitePath']}
    def update_existing(name):
        path=ROOT/name;raw=path.read_text('utf-8');updated=add_menu(raw)
        assert updated.count('data-teacher-menu')==2,(name,'Shared navigation must expose the directory twice')
        branch=None
        parts=name.split('/')
        if parts[0]=='지점안내' and len(parts)>=4:branch=by_slug.get(parts[2])
        elif parts[0] in ('전국학원','과목별학원'):
            summary=re.search(r'<section\b(?=[^>]*\bid="center-summary")[^>]*>[\s\S]*?</section>',raw)
            if summary:
                paths={unquote(urlsplit(unescape(u)).path) for u in re.findall(r'href="([^"]+)"',summary[0])}
                matches={u for u in paths if u in by_site}
                assert len(matches)<=1,(name,'Ambiguous branch links')
                if matches:branch=by_site[next(iter(matches))]
        if branch:
            if 'data-teacher-bridge=' not in updated:
                match=re.search(r'<section\b(?=[^>]*\bid="(?:center-summary|center-facts)")[^>]*>[\s\S]*?</section>',updated)
                if not match:raise ValueError('Cannot locate branch summary: '+name)
                updated=updated[:match.end()]+bridge(branch)+updated[match.end():]
        if updated!=raw:write(path,updated)
        return name,branch['slug'] if branch else None
    names=[name for name in baseline['files'] if name.endswith('.html')]
    with ThreadPoolExecutor(max_workers=12) as pool:
        for name,branch_name in pool.map(update_existing,names):
            nav_count+=1;changed.add(name)
            if branch_name:context_links[name]=branch_name
            if nav_count%1000==0:print(json.dumps({'navigationPages':nav_count,'total':len(names)}),flush=True)
    config=load(ROOT/'seo-descriptions.json')
    targets=[(None,hub(model,header,footer))]+[(b,branch_page(b,model,header,footer)) for b in model['branches']]
    for branch,(raw,desc) in targets:
        public=route(branch).lstrip('/')+'index.html';write(ROOT/public,raw);new_pages.append(public);changed.add(public)
        config['pages'][route(branch).rstrip('/')]={'description':desc,'sources':[desc]}
    write(ROOT/'seo-descriptions.json',json.dumps(config,ensure_ascii=False,indent=2))
    for photo in model['photos']:
        dest=ROOT/photo['path'];dest.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(teacher_dir/photo['source'],dest);changed.add(photo['path'])
    # Only pages with new main-content links receive a new sitemap date.
    sitemap=(out/'before-sitemap.xml').read_text('utf-8');context_paths={'/'+p.removesuffix('index.html') for p in context_links}
    def update_url(match):
        block=match[0];loc=re.search(r'<loc>(.*?)</loc>',block)[1]
        if unquote(urlsplit(unescape(loc)).path) not in context_paths:return block
        return re.sub(r'<lastmod>.*?</lastmod>',f'<lastmod>{DAY}</lastmod>',block) if '<lastmod>' in block else block.replace('</url>',f'<lastmod>{DAY}</lastmod></url>')
    sitemap=re.sub(r'<url\b[^>]*>[\s\S]*?</url>',update_url,sitemap)
    sitemap=sitemap.replace('</urlset>',''.join(f'<url><loc>{DOMAIN+H(route(b))}</loc><lastmod>{DAY}</lastmod></url>\n' for b in [None,*model['branches']])+'</urlset>')
    write(ROOT/'sitemap.xml',sitemap)
    changed|={'assets/teacher-directory.css','assets/teacher-directory.js','assets/unified-ui.css','sitemap.xml'}
    for name in sorted(changed):
        data=(ROOT/name).read_bytes();manifest['files'][name]=digest(data)
        if re.search(r'\.(html|css|js|xml|txt|json)$',name):manifest.setdefault('textSha256',{})[name]=digest(data.replace(b'\r\n',b'\n'))
    manifest['createdAt']=DAY;manifest['sitemapPages']=sum(n.endswith('.html') for n in manifest['files'])
    write(ROOT/'release-public-manifest.json',json.dumps(manifest,ensure_ascii=False,indent=2))
    result={'date':DAY,'teacherRows':model['sourceRows'],'sourceBranchLabels':205,'branches':len(model['branches']),'newPages':len(new_pages),'newPagePaths':new_pages,
      'matchedExistingBranches':193,'additionalBranches':11,'sharedImages':len(model['photos']),'navigationPages':nav_count,'contextLinks':context_links,'publicChanges':sorted(changed),
      'contentChangedPaths':sorted(context_paths),'sitemapPages':manifest['sitemapPages'],'publicFiles':len(manifest['files']),'sourceHash':model['sourceHash'],'deployed':False}
    write(out/'teacher-model.json',json.dumps(model,ensure_ascii=False,indent=2));write(out/'generation.json',json.dumps(result,ensure_ascii=False,indent=2))
    print(json.dumps({k:v for k,v in result.items() if k not in ('newPagePaths','contextLinks','publicChanges','contentChangedPaths')},ensure_ascii=False))
if __name__=='__main__':main()
