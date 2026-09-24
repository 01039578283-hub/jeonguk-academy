"""Verify the owner's sources and preserve existing content before adding a directory."""
from pathlib import Path
from collections import Counter
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import quote
from urllib.request import Request, urlopen
import csv, hashlib, json, re, shutil, subprocess, zipfile
import openpyxl
from lxml import html

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path('C:/Users/1992k/Desktop/센터정보')
REFERENCE = ROOT.parent / '새 홈페이지13'
DATA = ROOT / 'tools/data/branch-directory'
REPORT = ROOT / 'tools/reports/branch-directory-20260925'
ORIGIN = 'https://xn--3e0bl59bm0ad17a.com'
FAMILIES = ['과목별학원', '전국학원', '학습코칭', '학습가이드', '상담문의']

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def norm(s): return re.sub(r'[^가-힣A-Za-z0-9]', '', str(s))
def text(n): return ' '.join(n.text_content().split())
def save(p, v):
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(v, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
def pages(): return [ROOT/'index.html']+[p for f in FAMILIES for p in (ROOT/f).rglob('index.html')]

def main():
    REPORT.mkdir(parents=True, exist_ok=True)
    if not (REPORT/'baseline.json').exists():
        existing = pages()
        files = existing+[p for p in ROOT.iterdir() if p.is_file() and p.suffix in ['.xml','.txt','.json','.mjs'] and not p.name.startswith('.')]+[ROOT/'.vercelignore',ROOT/'.gitignore']+[p for p in (ROOT/'assets').iterdir() if p.suffix in ['.css','.js']]
        backup = Path('C:/Users/1992k/Desktop/CodexData/backups')/('jeonguk-before-branches-'+datetime.now().strftime('%Y%m%d-%H%M%S')+'.zip')
        backup.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(backup,'x',zipfile.ZIP_DEFLATED) as z:
            for p in files:z.write(p,p.relative_to(ROOT).as_posix())
        facts = {}
        for p in existing:
            raw=p.read_text('utf-8-sig');d=html.document_fromstring(raw)
            main=re.search(r'<main\b.*?</main>',raw,re.S)
            facts[p.relative_to(ROOT).as_posix()]={'sha256':sha(p),'mainSha256':hashlib.sha256(main[0].encode()).hexdigest() if main else None,'title':d.xpath('string(//title)'),'description':d.xpath('string(//meta[@name="description"]/@content)'),'canonical':d.xpath('string(//link[@rel="canonical"]/@href)'),'robots':d.xpath('string(//meta[@name="robots"]/@content)')}
        save(REPORT/'baseline.json',{'backup':str(backup),'pages':facts,'hashes':{p.relative_to(ROOT).as_posix():sha(p) for p in files},'head':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),'dirtyCount':len(subprocess.check_output(['git','status','--short'],cwd=ROOT,text=True,encoding='utf-8').splitlines())})

    source_audit=json.loads((REFERENCE/'tools/reports/smallclass-directory-20260924/source-audit.json').read_text('utf-8'))
    for s in source_audit['sources']:assert sha(SOURCE/s['file'])==s['sha256'], f'Source changed: {s["file"]}'
    data=json.loads((REFERENCE/'tools/data/smallclass-directory/centers.json').read_text('utf-8'))
    mapping=json.loads((REFERENCE/'tools/data/smallclass-directory/neighborhood-map.json').read_text('utf-8'))
    centers={c['key']:c for c in data['centers']}
    wb=openpyxl.load_workbook(SOURCE/'코칭센터_데이터_.xlsx',read_only=True,data_only=True)
    workbook_shapes=[{'sheet':s.title,'rows':s.max_row,'columns':s.max_column} for s in wb]
    ws=wb['센터정보']; rows=list(ws.values);columns=[str(c or '').replace('\n',' ') for c in rows[0]]
    exceptions=[]
    for c in centers.values():
        r=next((r for r in rows[1:] if r[0] and norm(r[0])==norm(c['key'])),None)
        if r is None:r=next((r for r in rows[1:] if r[11] and norm(r[11])==norm(c['address'])),None)
        if r is None:
            assert c['key']=='화성태안점',c['key']
            exceptions.append({'branch':c['key'],'basis':'최신 371개 코드의 주소·등록명·과목별 학년 사용. 원본 엑셀에 없는 운영 시간은 미확정으로 유지.'})
        else:assert norm(r[11])==norm(c['address']),c['key']
    wb.close()
    csvrows=list(csv.reader((SOURCE/'센터 정보 및 교육비 371개 코드_최신화.csv').open(encoding='utf-8-sig',newline='')))
    assert len(csvrows)==371 and all(len(r)==1 for r in csvrows)
    for m in mapping:
        c=centers[m['branch']]; doc=html.fromstring(csvrows[m['feeRow']-1][0]);plain=text(doc)
        assert norm(c['address']) in norm(plain),(c['key'],'address',m['feeRow'])
        assert norm(c['registeredName']) in norm(plain),(c['key'],'registration',m['feeRow'])
        for table in c['fees']:
            for row in table['rows']:
                for value in row:
                    assert norm(value) in norm(plain),(c['key'],'fee',value)
        assert not any(x in str(c['schools']) for x in ['[후보','확인 필요','휴·폐교','대상 미확인'])
    media={}
    for c in centers.values():
        for m in c['photos']+[c['map'],c['representative']]:
            src=REFERENCE/m['src'].lstrip('/');dst=ROOT/m['src'].lstrip('/')
            assert sha(src)==source_audit['media'][m['src']]['sha256']
            dst.parent.mkdir(parents=True,exist_ok=True)
            if not dst.exists():shutil.copy2(src,dst)
            assert sha(dst)==sha(src)
            media[m['src']]={'sha256':sha(dst),'bytes':dst.stat().st_size}
        c['reviewedAt']='2026-09-25'
    save(DATA/'centers.json',data)
    save(DATA/'neighborhood-map.json',mapping)
    result={'sources':source_audit['sources'],'workbookSheets':workbook_shapes,'workbookColumns':columns,'centers':len(centers),'regions':len(data['regions']),'mappedNeighborhoods':len(mapping),'mappedCenters':len(set(m['branch'] for m in mapping)),'feeModes':dict(Counter(c['feeMode'] for c in centers.values())),'photoModes':dict(Counter('center' if any(m['mode']=='center' for m in c['photos']) else 'common' for c in centers.values())),'exceptions':exceptions,'media':media,'unresolved':[]}
    save(REPORT/'source-audit.json',result)

    baseline=json.loads((REPORT/'baseline.json').read_text('utf-8'))
    sample=['index.html','전국학원/index.html','과목별학원/index.html','학습코칭/index.html','학습가이드/index.html']
    sample+=sorted(baseline['pages'])[::max(1,len(baseline['pages'])//15)]
    def check(rel):
        path='/'+rel.removesuffix('index.html') if rel!='index.html' else '/'
        with urlopen(Request(ORIGIN+quote(path,safe='/'),headers={'User-Agent':'Jeonguk-directory-QA/1.0'}),timeout=45) as res:
            raw=res.read().decode('utf-8');d=html.document_fromstring(raw)
            p=baseline['pages'][rel];desc=d.xpath('string(//meta[@name="description"]/@content)')
            return {'path':path,'status':res.status,'liveDescription':desc,'localDescription':p['description'],'descriptionMatches':desc==p['description'],'titleMatches':d.xpath('string(//title)')==p['title']}
    save(REPORT/'preexisting-live-comparison.json',list(ThreadPoolExecutor(5).map(check,list(dict.fromkeys(sample)))))
    print(json.dumps({k:v for k,v in result.items() if k not in ['media','sources','workbookColumns','workbookSheets']},ensure_ascii=False))
    print('Existing pages:',len(baseline['pages']),'Images:',len(media),'Backup:',baseline['backup'])
    home=html.document_fromstring((ROOT/'index.html').read_bytes())
    print('HEADER:',html.tostring(home.xpath('//header')[0],encoding='unicode'))
    print('FOOTER:',html.tostring(home.xpath('//footer')[0],encoding='unicode'))
    print('CSS:',home.xpath('//link[@rel="stylesheet"]/@href'))

if __name__=='__main__':main()
